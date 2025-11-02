# features/title.py
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from groq import Groq
from features.api import GROQ_API_KEY
from features.ompt import TITLE_PAGE_PROMPT
import json

client = Groq(api_key=GROQ_API_KEY)

def _add_section_break(paragraph):
    p = paragraph._element
    sectPr = OxmlElement("w:sectPr")
    type_el = OxmlElement("w:type")
    type_el.set(qn("w:val"), "nextPage")
    sectPr.append(type_el)
    p.addnext(sectPr)

def _classify_title(input_text: str):
    """
    Ask LLaMA/Groq for the title only.
    Returns {'title': ...} or empty dict on failure.
    """
    if not input_text.strip():
        return {}
    prompt = TITLE_PAGE_PROMPT.format(input_text=input_text)
    try:
        resp = client.chat.completions.create(
            model="llama-3.1-8b-instant",
            messages=[{"role": "user", "content": prompt}],
            max_tokens=100,
            temperature=0
        )
        payload = resp.choices[0].message.content
        data = json.loads(payload)
        return {"title": data.get("title", "")}
    except Exception:
        return {}

def fix_title_page(doc: Document, font_name="Times New Roman", font_size=12, max_title_paras=10):
    """
    Replace first max_title_paras paragraphs with APA-compliant title page.
    Returns (doc, title_text) — only the title for running head.
    """
    raw_candidates = [p.text.strip() for p in doc.paragraphs[:max_title_paras] if p.text.strip()]
    input_text = "\n".join(raw_candidates)

    # Only extract title
    data = _classify_title(input_text)
    if not data:
        data = {"title": raw_candidates[0] if raw_candidates else ""}

    # Remove old title paragraphs
    remove_count = min(max_title_paras, len(doc.paragraphs))
    for _ in range(remove_count):
        try:
            el = doc.paragraphs[0]._element
            el.getparent().remove(el)
        except Exception:
            break

    # Prepare APA title page lines
    nbsp = "\u00A0"
    title_lines = [
        (nbsp, False),
        (nbsp, False),
        (nbsp, False),
        (data.get("title", ""), True),
        (nbsp, False),
        (raw_candidates[1] if len(raw_candidates) > 1 else "", False),  # author
        (raw_candidates[2] if len(raw_candidates) > 2 else "", False),  # affiliation
        (raw_candidates[3] if len(raw_candidates) > 3 else "", False),  # course
        (raw_candidates[4] if len(raw_candidates) > 4 else "", False),  # instructor
        (raw_candidates[5] if len(raw_candidates) > 5 else "", False),  # date
    ]

    def _insert_before_first(text: str, bold: bool):
        if doc.paragraphs:
            ref = doc.paragraphs[0]
            p = ref.insert_paragraph_before(text, style="Normal")
        else:
            p = doc.add_paragraph(text)
        run = p.runs[0] if p.runs else p.add_run(text)
        run.bold = bool(bold)
        run.font.name = font_name
        run.font.size = Pt(font_size)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        return p

    for text, bold in reversed(title_lines):
        _insert_before_first(text, bold)

    # Section break after title page
    if doc.paragraphs:
        last_para = doc.paragraphs[len(title_lines) - 1]
        _add_section_break(last_para)

    # Return only the title for running head
    title_text = data.get("title", "")
    return doc, title_text
