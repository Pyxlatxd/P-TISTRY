# features/table.py
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def _clear_cell_tcBorders(cell):
    """Remove any tcBorders element from a cell (clears verticals/horizontals)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is not None:
        tcPr.remove(tcBorders)

def _set_cell_bottom_border(cell, sz='8', color='000000'):
    """Set bottom border on a cell (used for header-bottom and table-bottom)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    # remove existing bottom if present
    existing = tcBorders.find(qn('w:bottom'))
    if existing is not None:
        tcBorders.remove(existing)
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(sz))
    bottom.set(qn('w:space'), '0')
    bottom.set(qn('w:color'), color)
    tcBorders.append(bottom)

def _set_table_top_and_bottom(table, sz='8', color='000000'):
    """Set top and bottom borders at the table level (ensures top line is visible)."""
    tbl = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.append(tblPr)

    tblBorders = tblPr.find(qn('w:tblBorders'))
    if tblBorders is None:
        tblBorders = OxmlElement('w:tblBorders')
        tblPr.append(tblBorders)

    # remove any existing children
    for ch in list(tblBorders):
        tblBorders.remove(ch)

    top = OxmlElement('w:top')
    top.set(qn('w:val'), 'single')
    top.set(qn('w:sz'), str(sz))
    top.set(qn('w:space'), '0')
    top.set(qn('w:color'), color)
    tblBorders.append(top)

    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(sz))
    bottom.set(qn('w:space'), '0')
    bottom.set(qn('w:color'), color)
    tblBorders.append(bottom)

def append_apa_table(doc: Document, table_number: int, title: str, data: list,
                     font_name="Times New Roman", font_size=12):
    """
    Append an APA 7th edition style table to the end of an existing DOCX document.
    - Table-level top & bottom borders (ensures top line)
    - Header row cell-level bottom border (line under header)
    - No vertical borders
    - Double spacing applied
    Returns the modified doc.
    """
    if not data or not data[0]:
        raise ValueError("Data cannot be empty and must include headers.")

    # new page before table
    doc.add_page_break()

    # Table number
    num_para = doc.add_paragraph(f"Table {table_number}")
    num_run = num_para.runs[0] if num_para.runs else num_para.add_run(f"Table {table_number}")
    num_run.bold = True
    num_run.font.name = font_name
    num_run.font.size = Pt(font_size)
    num_para.paragraph_format.space_after = Pt(6)
    num_para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

    # Title (italic)
    title_para = doc.add_paragraph(title)
    title_run = title_para.runs[0] if title_para.runs else title_para.add_run(title)
    title_run.italic = True
    title_run.font.name = font_name
    title_run.font.size = Pt(font_size)
    title_para.paragraph_format.space_after = Pt(12)
    title_para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

    # Create table
    table = doc.add_table(rows=1, cols=len(data[0]))
    table.autofit = True

    # Fill header row (and clear any per-cell borders)
    for i, val in enumerate(data[0]):
        cell = table.rows[0].cells[i]
        _clear_cell_tcBorders(cell)
        para = cell.paragraphs[0]
        para.text = str(val)
        para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        run = para.runs[0]
        run.font.name = font_name
        run.font.size = Pt(font_size)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        run.bold = True  # optional: header bold for clarity

    # Add data rows, clearing cell borders
    for row_vals in data[1:]:
        new_row = table.add_row()
        for i, val in enumerate(row_vals):
            cell = new_row.cells[i]
            _clear_cell_tcBorders(cell)
            para = cell.paragraphs[0]
            para.text = str(val)
            para.alignment = (WD_PARAGRAPH_ALIGNMENT.RIGHT if isinstance(val, (int, float))
                              else WD_PARAGRAPH_ALIGNMENT.LEFT)
            run = para.runs[0]
            run.font.name = font_name
            run.font.size = Pt(font_size)
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

    # Set table-level top & bottom borders (ensures a visible top line)
    _set_table_top_and_bottom(table, sz='8', color='000000')

    # Set bottom border on header row (line under header)
    for cell in table.rows[0].cells:
        _set_cell_bottom_border(cell, sz='8', color='000000')

    # Ensure bottom border on last row (table-bottom)
    for cell in table.rows[-1].cells:
        _set_cell_bottom_border(cell, sz='8', color='000000')

    # spacing after table
    doc.add_paragraph()
    return doc
