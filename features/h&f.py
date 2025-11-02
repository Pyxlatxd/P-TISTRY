from docx import Document  
from docx.oxml import OxmlElement  
from docx.oxml.ns import qn  
  
def add_page_number_to_headers(doc):  
    for section in doc.sections:  
        # Access the header of the current section  
        header = section.header  
  
        # Create or access the first paragraph in the header  
        paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()  
        paragraph.alignment = 2  # Right alignment  
  
        # Add the page number field  
        run = paragraph.add_run()  
        fldChar1 = OxmlElement('w:fldChar')  # creates a new element  
        fldChar1.set(qn('w:fldCharType'), 'begin')  # sets attribute on element  
        instrText = OxmlElement('w:instrText')  
        instrText.set(qn('xml:space'), 'preserve')  # sets attribute on element  
        instrText.text = "PAGE"  
        fldChar2 = OxmlElement('w:fldChar')  
        fldChar2.set(qn('w:fldCharType'), 'end')  
        run._r.append(fldChar1)  
        run._r.append(instrText)  
        run._r.append(fldChar2)  
  
# Load your document  
doc = Document('example.docx')  
  
# Add page numbers to the top right of headers for all sections  
add_page_number_to_headers(doc)  
  
# Save the document with a new name  
doc.save('example_with_page_numbers.docx')  
