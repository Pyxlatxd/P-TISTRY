import json
import re
import inflect
from word2number import w2n
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK
from docx.text.paragraph import Paragraph
from groq import Groq
from features.api import GROQ_API_KEY

p = inflect.engine()
client = Groq(api_key=GROQ_API_KEY)

measurement_units = {"mg", "cm", "ml", "kg", "m", "km", "L", "g", "mm", "in", "ft", "%"}
time_units = {"second", "seconds", "minute", "minutes", "hour", "hours", "day", "days",
              "month", "months", "year", "years", "decade", "decades"}
all_units = measurement_units.union(time_units)

currency_mapping = {
    "dollar": "$", "dollars": "$",
    "euro": "€", "euros": "€",
    "pound": "£", "pounds": "£",
    "yen": "¥", "yenes": "¥",
    "rupee": "₹", "rupees": "₹"
}

def convert_word_to_number(word):
    try:
        return str(w2n.word_to_num(word))
    except ValueError:
        return word

def convert_numbers_in_text(text: str) -> str:
    words = text.split()
    processed = []
    i = 0
    while i < len(words):
        word = words[i]
        lower = word.lower().strip(",.!?")
        if lower in currency_mapping and i > 0:
            prev = processed[-1]
            try:
                num = str(w2n.word_to_num(prev.lower()))
            except Exception:
                num = prev if re.fullmatch(r'\d+(\.\d+)?', prev) else None
            if num:
                processed[-1] = f"{currency_mapping[lower]}{num}"
                i += 1
                continue
        elif lower in all_units and i > 0:
            prev = processed[-1]
            try:
                num = str(w2n.word_to_num(prev.lower()))
            except Exception:
                num = prev
            processed[-1] = f"{num} {lower}"
            i += 1
            continue
        elif lower.isalpha():
            processed.append(convert_word_to_number(lower))
        else:
            processed.append(word)
        i += 1
    return " ".join(processed)

def apply_apa_numerals(doc: Document, font_name="Times New Roman", font_size=12):
    """
    Let LLaMA decide which numbers to convert per APA rules,
    and append JSON reasoning at the end of the document.
    """

    # --- Gather paragraphs (after title page) ---
    paragraphs = []
    page_break_found = False

    for para in doc.paragraphs:
        if not isinstance(para, Paragraph):
            continue

        if any(getattr(run, "break_type", None) == WD_BREAK.PAGE for run in para.runs):
            page_break_found = True
            continue

        if not page_break_found:
            continue  # Skip title page

        if para.text.strip():
            paragraphs.append(para.text)

    if not paragraphs:
        return doc

    # --- LLaMA prompt ---
    prompt = (
        "You are an APA formatting assistant.\n"
        "Determine for each paragraph whether numerical expressions should be written as words or numerals.\n"
        "Base it on APA 7th edition rules (e.g., use numerals for 10 and above, exact units, time, ages, measurements, etc.).\n"
        "If it should not be reformatted, explain why.\n"
        "Return ONLY valid JSON like this:\n"
        '[{"text": "...", "convert": true, "reason": "APA requires numerals for 10 and above."}, ...]\n\n'
        f"Paragraphs to analyze:\n{json.dumps(paragraphs, ensure_ascii=False)}"
    )

    try:
        completion = client.chat.completions.create(
            model="llama-3.1-8b-instant",
            messages=[
                {"role": "system", "content": "You are an APA formatting assistant."},
                {"role": "user", "content": prompt}
            ],
            temperature=0,
            max_tokens=3000
        )

        raw = completion.choices[0].message.get("content") if isinstance(completion.choices[0].message, dict) \
            else completion.choices[0].message.content

        try:
            decisions = json.loads(raw.strip())
        except Exception as e:
            print(f"❌ JSON parse error: {e}")
            print(f"Raw output: {raw[:500]}")
            return doc

        # --- Apply LLaMA-guided numeral changes ---
        for para in doc.paragraphs:
            for item in decisions:
                if para.text.strip() and para.text.strip() == item.get("text", "").strip():
                    if item.get("convert"):
                        new_text = convert_numbers_in_text(para.text)
                        if new_text != para.text:
                            for run in list(para.runs):
                                para._element.remove(run._element)
                            new_run = para.add_run(new_text)
                            new_run.font.name = font_name
                            new_run.font.size = Pt(font_size)
                    break

        # --- Append reasoning JSON at the very end ---
        last_para = doc.add_paragraph()
        last_para.add_run("\n\nAPA Numeral Conversion Reasoning (AI Analysis):\n")
        last_para.add_run(json.dumps(decisions, indent=2, ensure_ascii=False))

        return doc

    except Exception as e:
        print(f"[Numeral Detection Fatal Error] {e}")
        return doc
