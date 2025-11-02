import json
from difflib import SequenceMatcher
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from groq import Groq
from features.api import GROQ_API_KEY

client = Groq(api_key=GROQ_API_KEY)

# --- Helper functions ---

def is_match(a, b):
    """Fuzzy match paragraphs ignoring small differences."""
    return SequenceMatcher(None, a.lower(), b.lower()).ratio() > 0.9

def apply_format(para, fmt):
    """Apply heading formatting to a docx paragraph."""
    alignment_map = {
        "LEFT": WD_ALIGN_PARAGRAPH.LEFT,
        "CENTER": WD_ALIGN_PARAGRAPH.CENTER,
        "RIGHT": WD_ALIGN_PARAGRAPH.RIGHT
    }
    para.alignment = alignment_map.get(fmt.get("alignment", "LEFT").upper(), WD_ALIGN_PARAGRAPH.LEFT)
    para.paragraph_format.first_line_indent = Inches(fmt.get("indent", 0))

    for run in para.runs:
        run.font.bold = fmt.get("bold", False)
        run.font.italic = fmt.get("italic", False)
        run.font.name = fmt.get("font", "Times New Roman")
        run.font.size = Pt(fmt.get("size", 12))

# --- Main function ---

def ai_detect_and_apply_headings(doc: Document, font_name="Times New Roman", font_size=12):
    """
    Detect APA heading levels (1–5) using LLaMA via Groq and apply formatting.
    Appends a JSON explanation at the very end of the document.
    """

    # --- Extract paragraph data ---
    doc_paragraphs = []
    blank_lines = 0

    for idx, para in enumerate(doc.paragraphs):
        text = para.text
        if not text.strip():
            blank_lines += 1
            continue

        run = para.runs[0] if para.runs else None
        left_indent = para.paragraph_format.left_indent
        left_indent_inches = left_indent.inches if left_indent else 0

        doc_paragraphs.append({
            "text": text,
            "bold": run.bold if run else False,
            "italic": run.italic if run else False,
            "underline": run.underline if run else False,
            "font": run.font.name if run else "Times New Roman",
            "size": run.font.size.pt if run and run.font.size else 12,
            "leading_blank_lines": blank_lines,
            "index": idx,
            "left_indent_inches": left_indent_inches
        })
        blank_lines = 0

    if not doc_paragraphs:
        return doc

    # --- Prepare prompt for LLaMA with reasoning ---
    prompt = (
        "You are an APA formatting assistant.\n"
        "Analyze each paragraph and determine intelligently its APA heading level (1–5) or null if not a heading.\n"
        "For each paragraph, provide your own comprehensive and insightful reasoning for why you chose this level heading for this paragraph. This is for better understanding on your model's decision-making.\n"
        "Return ONLY a JSON array like this:\n"
        '[{"text": "...", "level": 1, "reason": "Text is centered, bold, title-case..."}, ...]\n\n'
        "APA heading rules:\n"
        "1: Centered, Bold, Title Case\n"
        "2: Left-aligned, Bold, Title Case\n"
        "3: Left-aligned, Bold Italic, Title Case\n"
        "4: Indented, Bold, ends with a period\n"
        "5: Indented, Bold Italic, ends with a period\n\n"
        f"Paragraphs to analyze:\n{json.dumps(doc_paragraphs, ensure_ascii=False)}"
    )

    # --- Query model ---
    try:
        completion = client.chat.completions.create(
            model="llama-3.1-8b-instant",
            messages=[
                {"role": "system", "content": "You are an APA formatting assistant."},
                {"role": "user", "content": prompt}
            ],
            max_tokens=3000,
            temperature=0
        )

        raw = completion.choices[0].message.get("content") if isinstance(completion.choices[0].message, dict) \
            else completion.choices[0].message.content

        # --- Parse JSON response ---
        try:
            llama_response = json.loads(raw.strip())
        except Exception as e:
            print(f"❌ JSON parse error: {e}")
            print(f"Raw output (first 500 chars): {raw[:500]}")
            return doc

        # --- Heading format mapping ---
        heading_format = {
            1: {"bold": True, "italic": False, "alignment": "CENTER", "size": font_size, "indent": 0, "font": font_name},
            2: {"bold": True, "italic": False, "alignment": "LEFT",   "size": font_size, "indent": 0, "font": font_name},
            3: {"bold": True, "italic": True,  "alignment": "LEFT",   "size": font_size, "indent": 0, "font": font_name},
            4: {"bold": True, "italic": False, "alignment": "LEFT",   "size": font_size, "indent": 0.5, "font": font_name},
            5: {"bold": True, "italic": True,  "alignment": "LEFT",   "size": font_size, "indent": 0.5, "font": font_name},
        }

        # --- Apply formatting ---
        for para in doc.paragraphs:
            para_text = para.text
            for item in llama_response:
                if is_match(para_text, item.get("text", "")):
                    level = item.get("level")
                    if level is not None:
                        try:
                            lvl = int(level)
                            fmt = heading_format.get(lvl)
                            if fmt:
                                apply_format(para, fmt)
                        except (ValueError, KeyError):
                            continue

        # --- Append JSON reasoning at the very end ---
        last_para = doc.paragraphs[-1] if doc.paragraphs else doc.add_paragraph()
        last_para.add_run("\n\nAPA Heading Level Explanations (AI reasoning):\n")
        last_para.add_run(json.dumps(llama_response, indent=2, ensure_ascii=False))

        return doc

    except Exception as e:
        print(f"[Heading Detection Fatal Error] {e}")
        return doc
