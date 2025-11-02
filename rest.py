import streamlit as st
from docx import Document
from groq import Groq
from features.api import GROQ_API_KEY
import io

# Initialize Groq client
client = Groq(api_key=GROQ_API_KEY)

st.title("Groq LLaMA 3 to Word")

# File uploader for a blank Word doc
uploaded_file = st.file_uploader("Upload a blank Word document", type=["docx"])

if uploaded_file is not None:
    # Load the uploaded docx
    doc = Document(uploaded_file)

    # Run Groq API (classic prompt)
    chat_completion = client.chat.completions.create(
        messages=[
            {
                "role": "user",
                "content": "Explain the importance of fast language models",
            }
        ],
        model="llama-3.1-8b-instant",
        stream=False,
    )

    # Extract AI response
    response_text = chat_completion.choices[0].message.content

    # Insert response into doc
    doc.add_heading("Groq LLaMA 3 Response", level=1)
    doc.add_paragraph(response_text)

    # Save into memory (not disk)
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)

    # Download button
    st.download_button(
        label="Download Updated Word Document",
        data=output,
        file_name="groq_llama3_output.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
